"""Lo que ve la Administradora en los reportes: el Excel del RH1 respeta la sede
elegida, los grupos de residuo salen con su nombre y el dinero se lee como dinero."""
from datetime import date, time
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook

from core.models import AreaServicio, GestorExterno, Sede
from movimientos.models import Pesaje, TipoMovimiento
from residuos.models import CategoriaResiduo, DetalleResiduo, EntregaGestor

Usuario = get_user_model()
pytestmark = pytest.mark.django_db


@pytest.fixture
def administradora(db):
    return Usuario.objects.create_user(username="jefa", password="x", rol=Usuario.Rol.ADMIN)


@pytest.fixture
def dos_sedes_con_residuos(crear_movimiento):
    bio = CategoriaResiduo.objects.get(nombre="Biosanitarios")
    centro = Sede.objects.get(nombre="Centro de Cuidados")
    clinica = Sede.objects.get(nombre="Clínica Pabón")
    area_centro = AreaServicio.objects.get(sede=centro, nombre="UCI Coronaria")
    area_clinica = AreaServicio.objects.get(sede=clinica, nombre="Hemodinamia")
    for sede, area, kg in ((centro, area_centro, "4.00"), (clinica, area_clinica, "7.00")):
        mov = crear_movimiento(
            tipo=TipoMovimiento.RESIDUO_GENERACION, fecha=date(2026, 3, 10), hora=time(9, 0),
            sede=sede, area_origen=area,
        )
        DetalleResiduo.objects.create(movimiento=mov, categoria_residuo=bio, peso_kg=Decimal(kg))
    return centro, clinica


def _total_mes(resp):
    ws = load_workbook(BytesIO(resp.getvalue())).active
    return ws.cell(row=33, column=ws.max_column).value


# --- Excel del RH1 con sede -------------------------------------------------

def test_excel_rh1_sin_sede_suma_todas(client, administradora, dos_sedes_con_residuos):
    client.force_login(administradora)
    resp = client.get(reverse("reportes:exportar_rh1"), {"mes": "2026-03"})
    assert _total_mes(resp) == 11
    assert 'filename="rh1_2026-03.xlsx"' in resp["Content-Disposition"]


def test_excel_rh1_con_sede_solo_esa_sede(client, administradora, dos_sedes_con_residuos):
    centro, clinica = dos_sedes_con_residuos
    client.force_login(administradora)
    resp = client.get(reverse("reportes:exportar_rh1"), {"mes": "2026-03", "sede": centro.pk})
    assert _total_mes(resp) == 4
    assert "rh1_2026-03_centro-de-cuidados.xlsx" in resp["Content-Disposition"]
    resp = client.get(reverse("reportes:exportar_rh1"), {"mes": "2026-03", "sede": clinica.pk})
    assert _total_mes(resp) == 7


def test_excel_rh1_por_la_api_tambien_respeta_la_sede(api_client, administradora, dos_sedes_con_residuos):
    centro, _ = dos_sedes_con_residuos
    # Las exportaciones son vistas Django (con sesión), montadas también bajo /api/.
    api_client.force_login(administradora)
    resp = api_client.get(reverse("api-rh1-exportar"), {"mes": "2026-03", "sede": centro.pk})
    assert resp.status_code == 200
    assert _total_mes(resp) == 4


def test_sede_que_no_es_numero_se_ignora_sin_romper(client, api_client, administradora, dos_sedes_con_residuos):
    client.force_login(administradora)
    resp = client.get(reverse("reportes:exportar_rh1"), {"mes": "2026-03", "sede": "abc"})
    assert resp.status_code == 200
    assert _total_mes(resp) == 11
    api_client.force_authenticate(user=administradora)
    resp = api_client.get(reverse("api-rh1"), {"mes": "2026-03", "sede": "abc"})
    assert resp.status_code == 200
    assert Decimal(str(resp.data["total_mes"])) == Decimal("11")


# --- Nombres de grupo -------------------------------------------------------

def test_excel_de_consolidado_muestra_el_nombre_del_grupo(client, administradora, crear_movimiento):
    aprovechables = CategoriaResiduo.objects.get(nombre="Aprovechables")
    mov = crear_movimiento(tipo=TipoMovimiento.RESIDUO_GENERACION, fecha=date(2026, 3, 10), hora=time(9, 0))
    DetalleResiduo.objects.create(movimiento=mov, categoria_residuo=aprovechables, peso_kg=Decimal("3.00"))
    client.force_login(administradora)
    resp = client.get(reverse("reportes:exportar", args=["residuos_por_categoria"]))
    ws = load_workbook(BytesIO(resp.getvalue())).active
    grupos = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert "No peligroso" in grupos
    assert "NO_PELIGROSO" not in grupos


def test_pantalla_de_consolidados_muestra_el_nombre_del_grupo(client, administradora, crear_movimiento):
    aprovechables = CategoriaResiduo.objects.get(nombre="Aprovechables")
    mov = crear_movimiento(tipo=TipoMovimiento.RESIDUO_GENERACION, fecha=date(2026, 3, 10), hora=time(9, 0))
    DetalleResiduo.objects.create(movimiento=mov, categoria_residuo=aprovechables, peso_kg=Decimal("3.00"))
    client.force_login(administradora)
    cuerpo = client.get(reverse("reportes:consolidados")).content.decode()
    assert "No peligroso" in cuerpo
    assert "NO_PELIGROSO" not in cuerpo


# --- Dinero -----------------------------------------------------------------

def test_facturacion_muestra_pesos_con_signo_y_miles(client, administradora, crear_movimiento, usuario):
    gestor = GestorExterno.objects.create(nombre="SALVI S.A.S.", nit="900", tarifa_kg_vigente=Decimal("2500"))
    mov = crear_movimiento(
        tipo=TipoMovimiento.RESIDUO_RECOLECCION, fecha=date(2026, 3, 10), hora=time(9, 0),
        periodo_facturacion=date(2026, 3, 1),
    )
    Pesaje.objects.create(movimiento=mov, peso_total=Decimal("10.00"), tara=0, pesado_por=usuario)
    EntregaGestor.objects.create(
        movimiento=mov, gestor_externo=gestor, numero_factura="F1",
        kg_facturados=Decimal("10.00"), valor_facturado=Decimal("1234567.5"),
    )
    client.force_login(administradora)
    cuerpo = client.get(reverse("reportes:ambiental_facturacion"), {"mes": "2026-03"}).content.decode()
    assert "$ 1,234,567.50" in cuerpo
    assert "Total general facturado: $ 1,234,567.50" in cuerpo


def test_catalogo_muestra_la_tarifa_como_pesos(client, administradora):
    GestorExterno.objects.create(nombre="SALVI S.A.S.", nit="900", tarifa_kg_vigente=Decimal("2500"))
    client.force_login(administradora)
    assert "$ 2,500.00" in client.get(reverse("catalogos_parametros")).content.decode()
