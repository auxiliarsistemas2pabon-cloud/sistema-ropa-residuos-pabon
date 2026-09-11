from datetime import date, time
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook

from core.models import AreaServicio
from movimientos.models import Pesaje, TipoMovimiento
from reportes.services import por_jornada, ropa_por_sede, ropa_por_servicio
from residuos.models import CategoriaResiduo, DetalleResiduo

Usuario = get_user_model()
pytestmark = pytest.mark.django_db

DIA = date(2026, 3, 10)


@pytest.fixture
def administradora(db):
    return Usuario.objects.create_user(username="jefa", password="x", rol=Usuario.Rol.ADMIN)


@pytest.fixture
def datos(crear_movimiento, usuario, sede, area):
    quirofano, _ = AreaServicio.objects.get_or_create(
        sede=sede, nombre="Quirófano", defaults={"genera_ropa": True},
    )
    for a, kg, hora in [(area, "11.20", time(10, 0)), (quirofano, "23.40", time(10, 15)), (area, "6.00", time(14, 0))]:
        mov = crear_movimiento(tipo=TipoMovimiento.ROPA_SUCIA_ENTREGA, fecha=DIA, hora=hora, area_origen=a)
        Pesaje.objects.create(movimiento=mov, peso_total=Decimal(kg), tara=0, pesado_por=usuario)

    bio = CategoriaResiduo.objects.get(nombre="Biosanitarios")
    gen = crear_movimiento(tipo=TipoMovimiento.RESIDUO_GENERACION, fecha=DIA, hora=time(9, 0))
    DetalleResiduo.objects.create(movimiento=gen, categoria_residuo=bio, peso_kg=Decimal("4.00"))


def test_ropa_por_servicio_agrega(datos, sede):
    filas = ropa_por_servicio({"sede": sede})
    por_nombre = {f["movimiento__area_origen__nombre"]: f["kg"] for f in filas}
    assert por_nombre["Hemodinamia"] == Decimal("17.20")   # 11.20 + 6.00
    assert por_nombre["Quirófano"] == Decimal("23.40")


def test_ropa_por_sede_incluye_total(datos):
    filas, total = ropa_por_sede({})
    assert total == Decimal("40.60")


def test_por_jornada_separa_manana_y_tarde(datos):
    filas = {f["jornada"]: f for f in por_jornada({})}
    assert filas["Mañana"]["ropa_kg"] == Decimal("34.60")
    assert filas["Tarde"]["ropa_kg"] == Decimal("6.00")
    assert filas["Mañana"]["residuos_kg"] == Decimal("4.00")


def test_filtro_de_mes(datos):
    # el mes de marzo 2026 trae todo; abril, nada
    assert ropa_por_servicio({"desde": date(2026, 4, 1), "hasta": date(2026, 4, 30)}) == []


def test_consolidados_solo_administradora(client, usuario, administradora):
    client.force_login(usuario)
    assert client.get(reverse("reportes:consolidados")).status_code == 403

    client.force_login(administradora)
    assert client.get(reverse("reportes:consolidados")).status_code == 200


def test_exportar_devuelve_xlsx(client, administradora, datos):
    client.force_login(administradora)
    resp = client.get(reverse("reportes:exportar", args=["ropa_por_servicio"]))
    assert resp.status_code == 200
    assert resp["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "ropa_por_servicio.xlsx" in resp["Content-Disposition"]

    ws = load_workbook(BytesIO(resp.getvalue())).active
    assert [c.value for c in ws[1]] == ["Sede", "Servicio", "kg netos", "Movimientos"]
    valores = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(2, ws.max_row + 1)}
    assert valores["Hemodinamia"] == 17.2


def test_exportar_clave_desconocida_404(client, administradora):
    client.force_login(administradora)
    assert client.get(reverse("reportes:exportar", args=["inventado"])).status_code == 404
