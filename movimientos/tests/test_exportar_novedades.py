from datetime import date, time
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook

from .. import models as movimientos_models

Usuario = get_user_model()
pytestmark = pytest.mark.django_db


@pytest.fixture
def administradora(db):
    return Usuario.objects.create_user(username="jefa7", password="x", rol=Usuario.Rol.ADMIN)


@pytest.fixture
def dos_novedades(crear_movimiento, usuario):
    mov = crear_movimiento(
        tipo=movimientos_models.TipoMovimiento.ROPA_SUCIA_ENTREGA, fecha=date(2026, 3, 10), hora=time(10, 0),
    )
    movimientos_models.Novedad.objects.create(
        movimiento=mov, tipo_novedad=movimientos_models.TipoNovedad.ROPA_ROTA,
        observacion="Sábana rota", registrado_por=usuario,
    )
    movimientos_models.Novedad.objects.create(
        movimiento=mov, tipo_novedad=movimientos_models.TipoNovedad.DIFERENCIA_PESO,
        cantidad_afectada=2, observacion="", registrado_por=usuario,
    )
    return mov


def test_exportar_exige_administradora(client, usuario, administradora):
    client.force_login(usuario)
    assert client.get(reverse("movimientos:exportar_novedades")).status_code == 403

    client.force_login(administradora)
    assert client.get(reverse("movimientos:exportar_novedades")).status_code == 200


def test_exportar_devuelve_xlsx_con_las_novedades(client, administradora, dos_novedades):
    client.force_login(administradora)
    resp = client.get(reverse("movimientos:exportar_novedades"))
    assert resp["Content-Type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    ws = load_workbook(BytesIO(resp.getvalue())).active
    assert [c.value for c in ws[1]] == [
        "Fecha", "Tipo", "Sede", "Servicio", "kg afectados", "Observación", "Registró",
    ]
    observaciones = [ws.cell(row=r, column=6).value for r in range(2, ws.max_row + 1)]
    assert "Sábana rota" in observaciones


def test_exportar_respeta_el_filtro(client, administradora, dos_novedades):
    client.force_login(administradora)
    resp = client.get(
        reverse("movimientos:exportar_novedades"),
        {"tipo_novedad": movimientos_models.TipoNovedad.ROPA_ROTA},
    )
    ws = load_workbook(BytesIO(resp.getvalue())).active
    assert ws.max_row == 2  # cabecera + 1 fila
