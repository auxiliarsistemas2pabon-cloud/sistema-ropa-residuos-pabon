from datetime import date, datetime, time, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook

from core.models import GestorExterno
from movimientos.models import Movimiento, Pesaje, TipoMovimiento
from reportes.services import conciliacion_gestor, resumen_facturacion, rh1_del_mes
from residuos.models import CategoriaResiduo, ColumnaRH1, DetalleResiduo, EntregaGestor

Usuario = get_user_model()
pytestmark = pytest.mark.django_db


@pytest.fixture
def administradora(db):
    return Usuario.objects.create_user(username="jefa", password="x", rol=Usuario.Rol.ADMIN)


@pytest.fixture
def gestor(db):
    return GestorExterno.objects.create(
        nombre="SALVI S.A.S.", nit="900123456", tarifa_kg_vigente=Decimal("2500"),
    )


def test_columnas_rh1_sembradas():
    assert ColumnaRH1.objects.filter(activo=True).count() == 5
    assert ColumnaRH1.objects.filter(nombre="Biosanitarios").exists()


def test_rh1_del_mes_suma_por_dia(crear_movimiento, usuario):
    bio = CategoriaResiduo.objects.get(nombre="Biosanitarios")
    mov = crear_movimiento(tipo=TipoMovimiento.RESIDUO_GENERACION, fecha=date(2026, 3, 10), hora=time(9, 0))
    DetalleResiduo.objects.create(movimiento=mov, categoria_residuo=bio, peso_kg=Decimal("4.00"))

    datos = rh1_del_mes(2026, 3)
    assert len(datos["filas"]) == 31
    dia_10 = datos["filas"][9]
    col_bio = next(i for i, c in enumerate(datos["columnas"]) if c.nombre == "Biosanitarios")
    assert dia_10["celdas"][col_bio] == Decimal("4.00")
    assert datos["total_mes"] == Decimal("4.00")


def test_resumen_facturacion_separa_pendientes(crear_movimiento, usuario, gestor):
    def _entrega(periodo, kg, valor, creada_en):
        mov = crear_movimiento(
            tipo=TipoMovimiento.RESIDUO_RECOLECCION, fecha=periodo, hora=time(9, 0),
            periodo_facturacion=periodo,
        )
        Movimiento.objects.filter(pk=mov.pk).update(creado_en=creada_en)
        Pesaje.objects.create(movimiento=mov, peso_total=Decimal(kg), tara=0, pesado_por=usuario)
        return EntregaGestor.objects.create(
            movimiento=mov, gestor_externo=gestor, numero_factura="F1",
            kg_facturados=Decimal(kg), valor_facturado=Decimal(valor),
        )

    _entrega(date(2026, 3, 1), "10.00", "25000", datetime(2026, 3, 20, tzinfo=dt_timezone.utc))
    # del periodo de febrero pero cargada en marzo -> pendiente
    _entrega(date(2026, 2, 1), "6.00", "15000", datetime(2026, 3, 5, tzinfo=dt_timezone.utc))

    r = resumen_facturacion(2026, 3)
    assert r["total_actual"] == Decimal("25000")
    assert r["total_pendientes"] == Decimal("15000")
    assert r["total_general"] == Decimal("40000")


def test_conciliacion_marca_diferencia(crear_movimiento, usuario, gestor):
    mov = crear_movimiento(
        tipo=TipoMovimiento.RESIDUO_RECOLECCION, fecha=date(2026, 3, 10), hora=time(9, 0),
        periodo_facturacion=date(2026, 3, 1),
    )
    Pesaje.objects.create(movimiento=mov, peso_total=Decimal("10.30"), tara=Decimal("0.30"), pesado_por=usuario)
    EntregaGestor.objects.create(
        movimiento=mov, gestor_externo=gestor, numero_factura="F9",
        kg_facturados=Decimal("9.50"), valor_facturado=Decimal("1"),
    )

    filas = conciliacion_gestor(2026, 3)
    assert filas[0]["kg_interno"] == Decimal("10.00")
    assert filas[0]["kg_facturado"] == Decimal("9.50")
    assert filas[0]["diferencia"] == Decimal("0.50")


def test_exportar_rh1_xlsx(client, administradora, crear_movimiento, usuario):
    bio = CategoriaResiduo.objects.get(nombre="Biosanitarios")
    mov = crear_movimiento(tipo=TipoMovimiento.RESIDUO_GENERACION, fecha=date(2026, 3, 10), hora=time(9, 0))
    DetalleResiduo.objects.create(movimiento=mov, categoria_residuo=bio, peso_kg=Decimal("4.00"))

    client.force_login(administradora)
    resp = client.get(reverse("reportes:exportar_rh1"), {"mes": "2026-03"})
    assert resp.status_code == 200
    ws = load_workbook(BytesIO(resp.getvalue())).active
    assert ws["A1"].value == "Fecha"
    assert ws["A33"].value == "Total mes"  # 1 cabecera + 31 días + total


def test_ambiental_facturacion_solo_administradora(client, usuario, administradora):
    client.force_login(usuario)
    assert client.get(reverse("reportes:ambiental_facturacion")).status_code == 403
    client.force_login(administradora)
    assert client.get(reverse("reportes:ambiental_facturacion")).status_code == 200
