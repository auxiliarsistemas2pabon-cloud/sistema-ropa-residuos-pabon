from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

_CABECERA = Font(bold=True)
_DERECHA = Alignment(horizontal="right")


def libro_de_tabla(titulo, columnas, filas):
    """Workbook de openpyxl con una hoja: cabecera en negrita, ancho de
    columna automático y kilogramos alineados a la derecha."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]

    ws.append(columnas)
    for celda in ws[1]:
        celda.font = _CABECERA

    for fila in filas:
        ws.append(list(fila))

    for i, nombre in enumerate(columnas, start=1):
        letra = get_column_letter(i)
        largo = max([len(str(nombre))] + [len(str(f[i - 1])) for f in filas] + [10])
        ws.column_dimensions[letra].width = min(largo + 2, 60)
        if str(nombre).lower().startswith("kg"):
            for celda in ws[letra][1:]:
                celda.alignment = _DERECHA
                celda.number_format = "0.00"

    return wb
